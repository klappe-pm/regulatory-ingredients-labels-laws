#!/usr/bin/env python3
"""Render the regulatory-labeling prompt library manifests into a single Word document.

Reads public/prompts/<category>/.manifest.json for every category (in catalog
order), then emits public/prompts/regulatory-labeling-prompt-library.docx with
a hyperlinked table of contents, one prompt per page, and copy-ready blocks.
"""
import json
import re
from pathlib import Path

import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent.parent
PROMPTS_DIR = ROOT / "public" / "prompts"
OUT_PATH = PROMPTS_DIR / "regulatory-labeling-prompt-library.docx"

CATEGORY_ORDER = [
    "label-intake",
    "content-derivation",
    "claims-compliance",
    "review-and-verification",
    "version-reconciliation",
    "market-and-reporting",
    "communication-email",
    "durable-knowledge",
    "copilot-configuration",
]

GENERATED_DATE = "2026-07-22"


def bookmark_name(prompt_name):
    return re.sub(r"[^A-Za-z0-9_]", "_", prompt_name)[:40]


def add_bookmark(paragraph, name, bookmark_id):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_hyperlink(paragraph, anchor, text, color="0563C1", underline=True, size=None):
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    rpr.append(color_el)
    if underline:
        u_el = OxmlElement("w:u")
        u_el.set(qn("w:val"), "single")
        rpr.append(u_el)
    if size:
        sz_el = OxmlElement("w:sz")
        sz_el.set(qn("w:val"), str(size * 2))
        rpr.append(sz_el)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_toc_field(document):
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "Right-click and choose Update Field to build this native TOC."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r_element = run._r
    r_element.append(fld_begin)
    r_element.append(instr)
    r_element.append(fld_sep)
    r_element.append(fld_text)
    r_element.append(fld_end)
    return paragraph


def shaded_monospace_block(document, text):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.rows[0].cells[0]
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F2F2")
    tcPr.append(shd)
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "8")
        el.set(qn("w:color"), "999999")
        tcBorders.append(el)
    tcPr.append(tcBorders)
    cell.paragraphs[0].text = ""
    lines = text.split("\n")
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(line if line else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(9)
    return table


def labeled_line(document, label, value):
    p = document.add_paragraph()
    run_label = p.add_run(f"{label} ")
    run_label.bold = True
    p.add_run(value)
    return p


def return_to_toc(document):
    p = document.add_paragraph()
    add_internal_hyperlink(p, "TOC_TOP", "Return to table of contents", size=9)
    return p


def load_manifests():
    prompts = []
    for category in CATEGORY_ORDER:
        manifest_path = PROMPTS_DIR / category / ".manifest.json"
        with open(manifest_path) as f:
            items = json.load(f)
        for item in items:
            prompts.append(item)
    return prompts


def main():
    prompts = load_manifests()
    print(f"Loaded {len(prompts)} prompts across {len(CATEGORY_ORDER)} categories")

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Page 1: title, scope, usage note, TOC
    title_p = document.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_bookmark(title_p, "TOC_TOP", 0)
    run = title_p.add_run("Talking Rain, Regulatory Labeling")
    run.bold = True
    run.font.size = Pt(22)

    scope_p = document.add_paragraph()
    scope_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope_p.add_run(
        "Regulatory labeling prompt library for Microsoft 365 Copilot, "
        "beverage label compliance (US FDA jurisdiction)"
    ).italic = True

    gen_p = document.add_paragraph()
    gen_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gen_p.add_run(f"Generated: {GENERATED_DATE}")

    usage_p1 = document.add_paragraph()
    usage_p1.add_run(
        "Each entry below links to its prompt page. Every prompt page has a "
        "Return to table of contents link at top and bottom."
    )
    usage_p2 = document.add_paragraph()
    usage_p2.add_run(
        "Copy the shaded block on each prompt page directly into Microsoft 365 "
        "Copilot Chat (or Copilot in Outlook/Excel where noted)."
    )

    document.add_heading("Table of Contents", level=1)

    current_category = None
    for item in prompts:
        if item["category"] != current_category:
            current_category = item["category"]
            document.add_heading(current_category, level=2)
        p = document.add_paragraph(style="List Bullet")
        anchor = bookmark_name(item["name"])
        add_internal_hyperlink(p, anchor, item["name"])
        p.add_run(f": {item['when_to_use']}")

    document.add_heading("Native Word TOC (optional secondary aid)", level=2)
    document.add_paragraph(
        "The list above is the primary, hyperlinked table of contents. The "
        "field below is a native Word TOC field for readers who prefer it; "
        "update it via right-click > Update Field after opening the document."
    )
    add_toc_field(document)

    # One page per prompt
    bookmark_id = 1
    for item in prompts:
        document.add_page_break()
        return_to_toc(document)

        heading_p = document.add_paragraph()
        add_bookmark(heading_p, bookmark_name(item["name"]), bookmark_id)
        bookmark_id += 1
        run = heading_p.add_run(item["name"])
        run.bold = True
        run.font.size = Pt(16)

        labeled_line(document, "Description:", item["description"])
        labeled_line(document, "When to use:", item["when_to_use"])
        labeled_line(document, "What to use for:", item["what_to_use_for"])
        labeled_line(document, "Required inputs:", item["required_inputs"])
        labeled_line(document, "Expected outputs:", item["expected_outputs"])
        labeled_line(document, "Saves to:", f"{item['saves_to']} (visibility: {item['visibility']})")

        document.add_paragraph()
        shaded_monospace_block(document, item["prompt_text"])
        document.add_paragraph()

        return_to_toc(document)

    document.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
